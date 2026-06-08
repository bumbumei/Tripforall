"""Analyze the KTO Promton 2026 service intro PPTX template structure.

Lists per slide:
- slide layout name
- all shapes with their name, type, and current text content (if any)
- tables and their cell contents
"""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation
from pptx.util import Emu

TEMPLATE = "scripts/template.pptx"

prs = Presentation(TEMPLATE)
print(f"Total slides: {len(prs.slides)}\n")

for idx, slide in enumerate(prs.slides, 1):
    layout_name = slide.slide_layout.name
    print(f"{'=' * 70}")
    print(f"SLIDE {idx}  (layout: {layout_name})")
    print(f"{'=' * 70}")

    for shape_idx, shape in enumerate(slide.shapes):
        name = shape.name
        kind = type(shape).__name__
        has_tf = shape.has_text_frame
        is_table = shape.has_table
        is_pic = shape.shape_type == 13  # MSO_SHAPE_TYPE.PICTURE

        if is_table:
            tbl = shape.table
            print(f"  [{shape_idx}] {name} ({kind}) — TABLE {len(tbl.rows)}x{len(tbl.columns)}")
            for r_idx, row in enumerate(tbl.rows):
                for c_idx, cell in enumerate(row.cells):
                    txt = cell.text.strip().replace("\n", " | ")
                    if txt:
                        print(f"        row{r_idx}.col{c_idx}: {txt[:120]}")
        elif has_tf:
            txt = shape.text_frame.text.strip()
            preview = txt[:200].replace("\n", " ⏎ ") if txt else "(empty)"
            print(f"  [{shape_idx}] {name} ({kind}) — TEXT: {preview}")
        elif is_pic:
            print(f"  [{shape_idx}] {name} ({kind}) — PICTURE")
        else:
            print(f"  [{shape_idx}] {name} ({kind})")
    print()
