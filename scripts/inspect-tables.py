"""Deeper look at slide 4 table (9x2) and group shapes on slides 6/7."""

import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

from pptx import Presentation

prs = Presentation("scripts/template.pptx")

# Slide 4 — table 9x2
slide = prs.slides[3]  # 0-indexed
tbl = slide.shapes[3].table
print(f"SLIDE 4 — TABLE {len(tbl.rows)}x{len(tbl.columns)}")
for r_idx, row in enumerate(tbl.rows):
    for c_idx, cell in enumerate(row.cells):
        text = cell.text.strip().replace("\n", " | ")
        print(f"  row{r_idx}.col{c_idx}: '{text}'")
print()

# Slide 6/7 GroupShape — see what's inside
for slide_idx in [5, 6]:  # slides 6 and 7
    slide = prs.slides[slide_idx]
    print(f"SLIDE {slide_idx+1} group shapes")
    for shape in slide.shapes:
        if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
            print(f"  GroupShape {shape.name}")
            for inner in shape.shapes:
                txt = inner.text_frame.text.strip() if inner.has_text_frame else "(no text)"
                print(f"    inner: {inner.name} ({type(inner).__name__}) — {txt[:80]}")
